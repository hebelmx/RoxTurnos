#!/usr/bin/env python3
"""Git textconv driver: render an .xlsx workbook as plain text for diffs.

Emits one section per worksheet, with cells as tab-separated rows using
A1-style row numbers. Formulas are shown as `=FORMULA`; values otherwise.

Wire it up (per clone):
    git config diff.xlsx.textconv "python3 scripts/xlsx-textconv.py"
    git config diff.xlsx.binary true
"""
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    sys.stderr.write("xlsx-textconv: openpyxl not installed; pip install openpyxl\n")
    sys.exit(0)  # fall back to no diff rather than failing `git diff`


def render(path):
    wb = load_workbook(path, data_only=False, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# Sheet: {ws.title}  ({ws.max_row} rows x {ws.max_column} cols)")
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = [
                "" if v is None else str(v).replace("\t", " ").replace("\n", " ")
                for v in row
            ]
            if any(cells):
                out.append(f"{r}\t" + "\t".join(cells).rstrip("\t"))
        out.append("")
    return "\n".join(out)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.stderr.write("usage: xlsx-textconv.py <file.xlsx>\n")
        sys.exit(2)
    try:
        sys.stdout.write(render(sys.argv[1]))
    except Exception as exc:  # noqa: BLE001 - diff must never hard-fail
        sys.stderr.write(f"xlsx-textconv: {exc}\n")
        sys.exit(0)
